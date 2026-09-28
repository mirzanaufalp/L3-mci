#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
RESISTIVITY -> MCI ASSAY-CONSTRAINED
====================================

Tujuan:
1. Membaca data resistivity dari sheet Excel.
2. Membentuk MCI baseline 0-100 dari log-resistivity dengan clipping P5-P95.
3. Membaca data assay/lithology dari kolom "Zone Modifikasofti".
4. Mengubah zona menjadi target MCI:
      TP           -> 10
      soft         -> 30
      Hard cluster -> 65
      Hard         -> 90
5. Memberikan pengaruh constraint assay secara lokal menggunakan Gaussian
   berdasarkan jarak horizontal lintasan dan jarak vertikal.
6. Menghasilkan:
   - Excel hasil lengkap
   - PNG penampang MCI
   - CSV hasil (opsional, otomatis)

Catatan penting:
- MCI adalah indeks kompetensi RELATIF, bukan UCS/MPa.
- Perhitungan posisi assay mengasumsikan lubang bor vertikal:
      Z_mid = ZCollar - (From + To)/2
  Jika lubang miring, fungsi posisi assay perlu ditambah koreksi dip/azimuth.
- Parameter constraint harus dikalibrasi ulang bila data hardness/UCS/RQD tersedia.

Contoh:
    python L3.py L3.xlsx

Atau:
    python L3.py L3.xlsx --output-prefix L3_Hasil_MCI
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.tri as mtri

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter


# ============================================================
# 1. KONFIGURASI UTAMA
# ============================================================

RES_SHEET = "res_L3"
ASSAY_SHEET = "Assay L3"

# Kolom resistivity
COL_X = "X"
COL_Y = "Y"
COL_Z = "Z"
COL_RES = "RES"
COL_LINE = "LINE"

# Kolom assay
COL_BHID = "BHID"
COL_XCOLLAR = "XCollar"
COL_YCOLLAR = "YCollar"
COL_ZCOLLAR = "ZCollar"
COL_FROM = "From"
COL_TO = "To"
COL_ZONE = "Zone Modifikasofti"

# Normalisasi baseline resistivity
P_LOW = 5.0
P_HIGH = 95.0

# Target MCI berdasarkan constraint assay.
ZONE_MCI = {
    "tp": 10.0,
    "soft": 30.0,
    "hard cluster": 65.0,
    "hard": 90.0,
}

# Radius pengaruh assay.
H_SCALE_M = 15.0
V_SCALE_M = 0.40

# Maksimum kekuatan constraint assay.
MAX_ASSAY_WEIGHT = 0.95

# Bobot sangat kecil diabaikan agar constraint tidak memengaruhi seluruh lintasan.
MIN_EFFECTIVE_WEIGHT = 0.01

# Klasifikasi MCI akhir
MCI_CLASSES = [
    (0, 20, "Sangat rendah"),
    (20, 40, "Rendah"),
    (40, 60, "Sedang"),
    (60, 80, "Tinggi"),
    (80, 100.000001, "Sangat tinggi"),
]


# ============================================================
# 2. UTILITAS
# ============================================================

def normalize_name(value: object) -> str:
    """Normalisasi teks kategori zona."""
    if pd.isna(value):
        return ""
    return " ".join(str(value).strip().lower().split())


def require_columns(df: pd.DataFrame, columns: list[str], sheet_name: str) -> None:
    missing = [c for c in columns if c not in df.columns]
    if missing:
        raise ValueError(
            f"Sheet '{sheet_name}' tidak memiliki kolom: {missing}\n"
            f"Kolom yang tersedia: {list(df.columns)}"
        )


def classify_mci(value: float) -> str:
    if pd.isna(value):
        return ""
    value = float(np.clip(value, 0, 100))
    for lo, hi, label in MCI_CLASSES:
        if lo <= value < hi:
            return label
    return "Sangat tinggi"


def compute_profile_axis(x: np.ndarray, y: np.ndarray):
    """
    Membuat sumbu lintasan 1D dari koordinat XY dengan PCA/SVD.
    Cocok untuk lintasan ERT/GPR yang hampir linear.
    """
    xy = np.column_stack([x, y]).astype(float)
    center = np.nanmean(xy, axis=0)

    good = np.isfinite(xy).all(axis=1)
    if good.sum() < 2:
        raise ValueError("Koordinat XY valid terlalu sedikit untuk membentuk lintasan.")

    xy0 = xy[good] - center
    _, _, vh = np.linalg.svd(xy0, full_matrices=False)
    axis = vh[0].astype(float)

    # Orientasikan secara konsisten: dominan ke +X; bila X lemah, ke +Y.
    if abs(axis[0]) >= abs(axis[1]):
        if axis[0] < 0:
            axis *= -1
    else:
        if axis[1] < 0:
            axis *= -1

    s_raw = (xy - center) @ axis
    s0 = np.nanmin(s_raw)
    s = s_raw - s0

    return center, axis, s0, s


def project_to_profile(x, y, center, axis, s0):
    xy = np.column_stack([np.asarray(x, dtype=float), np.asarray(y, dtype=float)])
    return (xy - center) @ axis - s0


# ============================================================
# 3. MCI BASELINE DARI RESISTIVITY
# ============================================================

def compute_mci_resistivity(res: pd.Series):
    arr = pd.to_numeric(res, errors="coerce").to_numpy(dtype=float)
    valid = np.isfinite(arr) & (arr > 0)

    if valid.sum() < 5:
        raise ValueError("Nilai RES valid terlalu sedikit.")

    p_low = float(np.nanpercentile(arr[valid], P_LOW))
    p_high = float(np.nanpercentile(arr[valid], P_HIGH))

    if p_low <= 0 or p_high <= p_low:
        raise ValueError(f"P{P_LOW:g}/P{P_HIGH:g} RES tidak valid: {p_low}, {p_high}")

    mci = np.full(arr.shape, np.nan, dtype=float)
    log_low = math.log10(p_low)
    log_high = math.log10(p_high)

    mci[valid] = 100.0 * (
        (np.log10(arr[valid]) - log_low) / (log_high - log_low)
    )
    mci[valid] = np.clip(mci[valid], 0, 100)

    return mci, p_low, p_high


# ============================================================
# 4. PREPARASI ASSAY CONSTRAINT
# ============================================================

def prepare_assay(
    assay: pd.DataFrame,
    profile_center: np.ndarray,
    profile_axis: np.ndarray,
    profile_s0: float,
) -> pd.DataFrame:
    required = [
        COL_BHID, COL_XCOLLAR, COL_YCOLLAR, COL_ZCOLLAR,
        COL_FROM, COL_TO, COL_ZONE,
    ]
    require_columns(assay, required, ASSAY_SHEET)

    a = assay.copy()

    a["Zone_Normalized"] = a[COL_ZONE].map(normalize_name)
    a["MCI_Assay_Target"] = a["Zone_Normalized"].map(ZONE_MCI)

    for col in [COL_XCOLLAR, COL_YCOLLAR, COL_ZCOLLAR, COL_FROM, COL_TO]:
        a[col] = pd.to_numeric(a[col], errors="coerce")

    # Hanya zona yang memiliki target MCI digunakan sebagai constraint.
    a = a[
        a["MCI_Assay_Target"].notna()
        & a[[COL_XCOLLAR, COL_YCOLLAR, COL_ZCOLLAR, COL_FROM, COL_TO]].notna().all(axis=1)
    ].copy()

    if a.empty:
        raise ValueError(
            "Tidak ada data assay yang dapat digunakan sebagai constraint. "
            "Periksa Zone Modifikasofti dan ZONE_MCI."
        )

    a["Depth_Mid_m"] = (a[COL_FROM] + a[COL_TO]) / 2.0

    # ASUMSI lubang vertikal.
    a["Assay_Z_Mid"] = a[COL_ZCOLLAR] - a["Depth_Mid_m"]
    a["Assay_Z_From"] = a[COL_ZCOLLAR] - a[COL_FROM]
    a["Assay_Z_To"] = a[COL_ZCOLLAR] - a[COL_TO]

    a["Profile_m"] = project_to_profile(
        a[COL_XCOLLAR].to_numpy(),
        a[COL_YCOLLAR].to_numpy(),
        profile_center,
        profile_axis,
        profile_s0,
    )

    return a


# ============================================================
# 5. GABUNGKAN RESISTIVITY + ASSAY
# ============================================================

def apply_assay_constraint(
    res_df: pd.DataFrame,
    assay_df: pd.DataFrame,
) -> pd.DataFrame:
    out = res_df.copy()

    s_res = out["Profile_m"].to_numpy(dtype=float)
    z_res = pd.to_numeric(out[COL_Z], errors="coerce").to_numpy(dtype=float)
    base = out["MCI_RES_Base"].to_numpy(dtype=float)

    s_assay = assay_df["Profile_m"].to_numpy(dtype=float)
    z_assay = assay_df["Assay_Z_Mid"].to_numpy(dtype=float)
    target = assay_df["MCI_Assay_Target"].to_numpy(dtype=float)

    # Matrix distance: N_res x N_assay
    dh = s_res[:, None] - s_assay[None, :]
    dv = z_res[:, None] - z_assay[None, :]

    w = np.exp(
        -0.5 * (
            (dh / H_SCALE_M) ** 2
            + (dv / V_SCALE_M) ** 2
        )
    )

    w[w < MIN_EFFECTIVE_WEIGHT] = 0.0

    wsum = w.sum(axis=1)
    wmax = w.max(axis=1) if w.shape[1] else np.zeros(len(out))

    assay_target = np.full(len(out), np.nan, dtype=float)
    constrained = wsum > 0

    if constrained.any():
        assay_target[constrained] = (
            w[constrained] @ target
        ) / wsum[constrained]

    constraint_weight = MAX_ASSAY_WEIGHT * wmax
    constraint_weight = np.clip(constraint_weight, 0, MAX_ASSAY_WEIGHT)

    final = base.copy()
    usable = constrained & np.isfinite(base) & np.isfinite(assay_target)
    final[usable] = (
        (1.0 - constraint_weight[usable]) * base[usable]
        + constraint_weight[usable] * assay_target[usable]
    )
    final = np.clip(final, 0, 100)

    out["MCI_Assay_Target_Local"] = assay_target
    out["Constraint_Weight"] = constraint_weight
    out["MCI_AssayConstrained"] = final
    out["MCI_Class"] = [classify_mci(v) for v in final]

    return out


# ============================================================
# 6. PENAMPANG
# ============================================================

def plot_section(
    result: pd.DataFrame,
    assay: pd.DataFrame,
    output_png: Path,
    title: str,
):
    x = result["Profile_m"].to_numpy(dtype=float)
    z = pd.to_numeric(result[COL_Z], errors="coerce").to_numpy(dtype=float)
    mci = result["MCI_AssayConstrained"].to_numpy(dtype=float)

    good = np.isfinite(x) & np.isfinite(z) & np.isfinite(mci)
    xg, zg, mg = x[good], z[good], mci[good]

    fig, ax = plt.subplots(figsize=(14, 6))

    plotted_contour = False
    if len(xg) >= 3:
        try:
            triang = mtri.Triangulation(xg, zg)
            cntr = ax.tricontourf(
                triang,
                mg,
                levels=np.linspace(0, 100, 21),
            )
            plotted_contour = True
        except Exception:
            plotted_contour = False

    if not plotted_contour:
        cntr = ax.scatter(
            xg, zg,
            c=mg,
            s=18,
            vmin=0,
            vmax=100,
        )

    cbar = fig.colorbar(cntr, ax=ax)
    cbar.set_label("MCI Assay-Constrained (0-100)")

    for _, row in assay.iterrows():
        ax.plot(
            [row["Profile_m"], row["Profile_m"]],
            [row["Assay_Z_From"], row["Assay_Z_To"]],
            linewidth=4,
        )

    hole_info = (
        assay[[COL_BHID, "Profile_m", COL_ZCOLLAR]]
        .drop_duplicates(subset=[COL_BHID])
    )
    for _, row in hole_info.iterrows():
        ax.text(
            row["Profile_m"],
            row[COL_ZCOLLAR],
            str(row[COL_BHID]),
            rotation=90,
            va="bottom",
            ha="center",
            fontsize=8,
        )

    ax.set_xlabel("Jarak sepanjang lintasan (m)")
    ax.set_ylabel("Elevasi / Z (m)")
    ax.set_title(title)
    ax.grid(alpha=0.2)
    fig.tight_layout()
    fig.savefig(output_png, dpi=200, bbox_inches="tight")
    plt.close(fig)


# ============================================================
# 7. OUTPUT EXCEL
# ============================================================

def auto_format_excel(path: Path):
    """Formatting sederhana agar workbook hasil mudah dibaca."""
    wb = load_workbook(path)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        for col_cells in ws.columns:
            letter = get_column_letter(col_cells[0].column)
            max_len = 0
            for cell in col_cells[:200]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, len(value))
            ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 28)

    # Conditional format MCI pada sheet hasil
    if "Res_MCI" in wb.sheetnames:
        ws = wb["Res_MCI"]
        headers = {cell.value: cell.column for cell in ws[1]}
        mci_col = headers.get("MCI_AssayConstrained")
        if mci_col:
            letter = get_column_letter(mci_col)
            ws.conditional_formatting.add(
                f"{letter}2:{letter}{ws.max_row}",
                ColorScaleRule(
                    start_type="num", start_value=0, start_color="F8696B",
                    mid_type="num", mid_value=50, mid_color="FFEB84",
                    end_type="num", end_value=100, end_color="63BE7B",
                ),
            )

    wb.save(path)


def write_excel(
    output_xlsx: Path,
    result: pd.DataFrame,
    assay: pd.DataFrame,
    p_low: float,
    p_high: float,
):
    params = pd.DataFrame(
        [
            ["Metode baseline", "log10 resistivity P5-P95"],
            ["P low (%)", P_LOW],
            ["P high (%)", P_HIGH],
            [f"RES P{P_LOW:g} (ohm.m)", p_low],
            [f"RES P{P_HIGH:g} (ohm.m)", p_high],
            ["Horizontal scale constraint (m)", H_SCALE_M],
            ["Vertical scale constraint (m)", V_SCALE_M],
            ["Maximum assay weight", MAX_ASSAY_WEIGHT],
            ["Minimum effective Gaussian weight", MIN_EFFECTIVE_WEIGHT],
            ["TP target MCI", ZONE_MCI["tp"]],
            ["soft target MCI", ZONE_MCI["soft"]],
            ["Hard cluster target MCI", ZONE_MCI["hard cluster"]],
            ["Hard target MCI", ZONE_MCI["hard"]],
            [
                "Formula baseline",
                "100*clip[(log10(RES)-log10(P5))/(log10(P95)-log10(P5)),0,1]",
            ],
            [
                "Formula constraint",
                "MCI_final=(1-W)*MCI_RES + W*MCI_assay; W=max_assay_weight*max(Gaussian)",
            ],
            [
                "Asumsi posisi assay",
                "Lubang vertikal: Zmid=ZCollar-(From+To)/2",
            ],
        ],
        columns=["Parameter", "Value"],
    )

    assay_export_cols = [
        COL_BHID,
        COL_XCOLLAR,
        COL_YCOLLAR,
        COL_ZCOLLAR,
        COL_FROM,
        COL_TO,
        COL_ZONE,
        "Zone_Normalized",
        "MCI_Assay_Target",
        "Depth_Mid_m",
        "Assay_Z_From",
        "Assay_Z_Mid",
        "Assay_Z_To",
        "Profile_m",
    ]
    assay_out = assay[assay_export_cols].copy()

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        result.to_excel(writer, sheet_name="Res_MCI", index=False)
        assay_out.to_excel(writer, sheet_name="Assay_Constraint", index=False)
        params.to_excel(writer, sheet_name="MCI_Parameters", index=False)

    auto_format_excel(output_xlsx)


# ============================================================
# 8. MAIN PROCESS
# ============================================================

def run(input_xlsx: Path, output_prefix: str | None = None):
    if not input_xlsx.exists():
        raise FileNotFoundError(input_xlsx)

    if output_prefix is None:
        output_prefix = f"{input_xlsx.stem}_MCI_AssayConstrained"

    out_dir = input_xlsx.parent
    output_xlsx = out_dir / f"{output_prefix}.xlsx"
    output_png = out_dir / f"{output_prefix}_Section.png"
    output_csv = out_dir / f"{output_prefix}.csv"

    print(f"[1/6] Membaca: {input_xlsx}")
    res = pd.read_excel(input_xlsx, sheet_name=RES_SHEET)
    assay_raw = pd.read_excel(input_xlsx, sheet_name=ASSAY_SHEET)

    require_columns(
        res,
        [COL_X, COL_Y, COL_Z, COL_RES],
        RES_SHEET,
    )

    for col in [COL_X, COL_Y, COL_Z, COL_RES]:
        res[col] = pd.to_numeric(res[col], errors="coerce")

    print("[2/6] Membentuk jarak profil...")
    center, axis, s0, s_res = compute_profile_axis(
        res[COL_X].to_numpy(),
        res[COL_Y].to_numpy(),
    )
    res["Profile_m"] = s_res

    print("[3/6] Menghitung MCI baseline dari resistivity...")
    mci_base, p_low, p_high = compute_mci_resistivity(res[COL_RES])
    res["MCI_RES_Base"] = mci_base

    print(
        f"      RES P{P_LOW:g} = {p_low:.3f} ohm.m | "
        f"P{P_HIGH:g} = {p_high:.3f} ohm.m"
    )

    print("[4/6] Membentuk constraint assay...")
    assay = prepare_assay(
        assay_raw,
        center,
        axis,
        s0,
    )
    print(
        f"      {len(assay)} interval assay valid, "
        f"{assay[COL_BHID].nunique()} hole."
    )

    print("[5/6] Menggabungkan MCI resistivity + assay...")
    result = apply_assay_constraint(res, assay)

    print("[6/6] Menulis output Excel, CSV, dan penampang...")
    write_excel(output_xlsx, result, assay, p_low, p_high)
    result.to_csv(output_csv, index=False)

    plot_section(
        result,
        assay,
        output_png,
        title=f"{input_xlsx.stem} - MCI Resistivity dengan Constraint Assay",
    )

    print("\nSELESAI")
    print(f"Excel     : {output_xlsx}")
    print(f"CSV       : {output_csv}")
    print(f"Penampang : {output_png}")
    print(f"Median MCI final : {np.nanmedian(result['MCI_AssayConstrained']):.2f}")
    print(
        f"Proporsi titik yang terpengaruh assay "
        f"(W > {MIN_EFFECTIVE_WEIGHT:.2f}): "
        f"{100*np.mean(result['Constraint_Weight'] > MIN_EFFECTIVE_WEIGHT):.1f}%"
    )


def main():
    parser = argparse.ArgumentParser(
        description="Konversi resistivity menjadi MCI dengan Zone Modifikasofti sebagai constraint assay."
    )
    parser.add_argument(
        "input_xlsx",
        type=Path,
        help="File input Excel, contoh: L3.xlsx",
    )
    parser.add_argument(
        "--output-prefix",
        default=None,
        help="Prefix file hasil. Default: <nama_input>_MCI_AssayConstrained",
    )
    args = parser.parse_args()

    run(args.input_xlsx, args.output_prefix)


if __name__ == "__main__":
    main()
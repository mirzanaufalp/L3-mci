import streamlit as st
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import matplotlib.tri as mtri
import io
import math
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

# Set konfigurasi halaman web
st.set_page_config(page_title="MCI Assay-Constrained Processor", layout="wide")

st.title("⚡ Resistivity -> MCI Assay-Constrained Processor")
st.markdown("Aplikasi konversi data resistivity geolistrik menjadi Mass Competency Index (MCI) ber-constraint data bor/assay.")

# Sidebar untuk Pengaturan Parameter
st.sidebar.header("⚙️ Konfigurasi Parameter")
h_scale = st.sidebar.slider("Horizontal Scale Constraint (m)", 1.0, 50.0, 15.0, 0.5)
v_scale = st.sidebar.slider("Vertical Scale Constraint (m)", 0.1, 5.0, 0.4, 0.05)
max_weight = st.sidebar.slider("Max Assay Weight", 0.1, 1.0, 0.95, 0.05)

st.sidebar.subheader("🎯 Target MCI per Zona")
target_tp = st.sidebar.number_input("TP", value=10.0)
target_soft = st.sidebar.number_input("soft", value=30.0)
target_hard_cluster = st.sidebar.number_input("Hard cluster", value=65.0)
target_hard = st.sidebar.number_input("Hard", value=90.0)

ZONE_MCI = {
    "tp": target_tp,
    "soft": target_soft,
    "hard cluster": target_hard_cluster,
    "hard": target_hard,
}

P_LOW, P_HIGH = 5.0, 95.0
MIN_EFFECTIVE_WEIGHT = 0.01

# --- Fungsi Utility ---
def normalize_name(value):
    if pd.isna(value): return ""
    return " ".join(str(value).strip().lower().split())

def compute_profile_axis(x, y):
    xy = np.column_stack([x, y]).astype(float)
    center = np.nanmean(xy, axis=0)
    good = np.isfinite(xy).all(axis=1)
    if good.sum() < 2:
        raise ValueError("Koordinat XY valid terlalu sedikit.")
    xy0 = xy[good] - center
    _, _, vh = np.linalg.svd(xy0, full_matrices=False)
    axis = vh[0].astype(float)
    if abs(axis[0]) >= abs(axis[1]):
        if axis[0] < 0: axis *= -1
    else:
        if axis[1] < 0: axis *= -1
    s_raw = (xy - center) @ axis
    s0 = np.nanmin(s_raw)
    return center, axis, s0, s_raw - s0

def project_to_profile(x, y, center, axis, s0):
    xy = np.column_stack([np.asarray(x, dtype=float), np.asarray(y, dtype=float)])
    return (xy - center) @ axis - s0

def compute_mci_resistivity(res):
    arr = pd.to_numeric(res, errors="coerce").to_numpy(dtype=float)
    valid = np.isfinite(arr) & (arr > 0)
    if valid.sum() < 5: raise ValueError("Nilai RES valid terlalu sedikit.")
    p_low, p_high = float(np.nanpercentile(arr[valid], P_LOW)), float(np.nanpercentile(arr[valid], P_HIGH))
    mci = np.full(arr.shape, np.nan, dtype=float)
    log_low, log_high = math.log10(p_low), math.log10(p_high)
    mci[valid] = 100.0 * ((np.log10(arr[valid]) - log_low) / (log_high - log_low))
    return np.clip(mci, 0, 100), p_low, p_high

def prepare_assay(assay, profile_center, profile_axis, profile_s0):
    a = assay.copy()
    a["Zone_Normalized"] = a["Zone Modifikasofti"].map(normalize_name)
    a["MCI_Assay_Target"] = a["Zone_Normalized"].map(ZONE_MCI)
    for col in ["XCollar", "YCollar", "ZCollar", "From", "To"]:
        a[col] = pd.to_numeric(a[col], errors="coerce")
    a = a[a["MCI_Assay_Target"].notna() & a[["XCollar", "YCollar", "ZCollar", "From", "To"]].notna().all(axis=1)].copy()
    a["Depth_Mid_m"] = (a["From"] + a["To"]) / 2.0
    a["Assay_Z_Mid"] = a["ZCollar"] - a["Depth_Mid_m"]
    a["Assay_Z_From"] = a["ZCollar"] - a["From"]
    a["Assay_Z_To"] = a["ZCollar"] - a["To"]
    a["Profile_m"] = project_to_profile(a["XCollar"].to_numpy(), a["YCollar"].to_numpy(), profile_center, profile_axis, profile_s0)
    return a

def apply_assay_constraint(res_df, assay_df):
    out = res_df.copy()
    s_res = out["Profile_m"].to_numpy(dtype=float)
    z_res = pd.to_numeric(out["Z"], errors="coerce").to_numpy(dtype=float)
    base = out["MCI_RES_Base"].to_numpy(dtype=float)
    s_assay, z_assay = assay_df["Profile_m"].to_numpy(dtype=float), assay_df["Assay_Z_Mid"].to_numpy(dtype=float)
    target = assay_df["MCI_Assay_Target"].to_numpy(dtype=float)
    
    dh = s_res[:, None] - s_assay[None, :]
    dv = z_res[:, None] - z_assay[None, :]
    w = np.exp(-0.5 * ((dh / h_scale) ** 2 + (dv / v_scale) ** 2))
    w[w < MIN_EFFECTIVE_WEIGHT] = 0.0
    
    wsum, wmax = w.sum(axis=1), w.max(axis=1) if w.shape[1] else np.zeros(len(out))
    assay_target = np.full(len(out), np.nan, dtype=float)
    constrained = wsum > 0
    if constrained.any():
        assay_target[constrained] = (w[constrained] @ target) / wsum[constrained]
        
    constraint_weight = np.clip(max_weight * wmax, 0, max_weight)
    final = base.copy()
    usable = constrained & np.isfinite(base) & np.isfinite(assay_target)
    final[usable] = (1.0 - constraint_weight[usable]) * base[usable] + constraint_weight[usable] * assay_target[usable]
    
    out["MCI_Assay_Target_Local"] = assay_target
    out["Constraint_Weight"] = constraint_weight
    out["MCI_AssayConstrained"] = np.clip(final, 0, 100)
    return out

# --- UI File Upload ---
uploaded_file = st.file_uploader("Upload File Excel (`L3.xlsx`)", type=["xlsx"])

if uploaded_file is not None:
    try:
        res = pd.read_excel(uploaded_file, sheet_name="res_L3")
        assay_raw = pd.read_excel(uploaded_file, sheet_name="Assay L3")
        
        for col in ["X", "Y", "Z", "RES"]:
            res[col] = pd.to_numeric(res[col], errors="coerce")
            
        center, axis, s0, s_res = compute_profile_axis(res["X"].to_numpy(), res["Y"].to_numpy())
        res["Profile_m"] = s_res
        
        mci_base, p_low, p_high = compute_mci_resistivity(res["RES"])
        res["MCI_RES_Base"] = mci_base
        
        assay = prepare_assay(assay_raw, center, axis, s0)
        result = apply_assay_constraint(res, assay)
        
        st.success("✅ Pemrosesan Data MCI Berhasil!")
        
        # Display Visual Plot
        st.subheader("🖼️ Penampang Profil MCI")
        fig, ax = plt.subplots(figsize=(12, 5))
        xg, zg, mg = result["Profile_m"].to_numpy(), result["Z"].to_numpy(), result["MCI_AssayConstrained"].to_numpy()
        good = np.isfinite(xg) & np.isfinite(zg) & np.isfinite(mg)
        
        triang = mtri.Triangulation(xg[good], zg[good])
        cntr = ax.tricontourf(triang, mg[good], levels=np.linspace(0, 100, 21))
        cbar = fig.colorbar(cntr, ax=ax)
        cbar.set_label("MCI Assay-Constrained (0-100)")
        
        for _, row in assay.iterrows():
            ax.plot([row["Profile_m"], row["Profile_m"]], [row["Assay_Z_From"], row["Assay_Z_To"]], linewidth=3, color="black")
            
        ax.set_xlabel("Jarak Profil (m)")
        ax.set_ylabel("Elevasi Z (m)")
        ax.grid(alpha=0.3)
        st.pyplot(fig)
        
        # Download Buttons
        st.subheader("📥 Download Hasil")
        csv_data = result.to_csv(index=False).encode('utf-8')
        st.download_button("Download CSV Hasil", data=csv_data, file_name="MCI_Hasil.csv", mime="text/csv")
        
    except Exception as e:
        st.error(f"Terjadi kesalahan saat memproses file: {e}")
